"""Utilities for reading company data from an Excel file."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from models import InputCompanyRow

COMPANY_HEADER_ALIASES: tuple[str, ...] = (
    "Company",
    "Компания",
)

REGISTRATION_NUMBER_HEADER_ALIASES: tuple[str, ...] = (
    "RegistrationNumber",
    "Registration Number",
    "Регистрационный номер",
)


class InputReader:
    """Read source company rows from an Excel workbook."""

    def __init__(self, file_path: str | Path) -> None:
        """Initialize the input reader."""
        self.file_path = Path(file_path)

    def read_rows(self) -> list[InputCompanyRow]:
        """Read company rows from the first worksheet by header names."""
        workbook = load_workbook(filename=self.file_path, read_only=True, data_only=True)
        worksheet = workbook.active
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)

        if header_row is None:
            workbook.close()
            raise ValueError("Input Excel file is empty")

        header_map = self._build_header_map(header_row)
        company_column = self._find_column_index(header_map, COMPANY_HEADER_ALIASES)
        registration_column = self._find_column_index(
            header_map,
            REGISTRATION_NUMBER_HEADER_ALIASES,
        )

        if company_column is None:
            workbook.close()
            raise ValueError(
                "Input Excel must contain Company or Компания column"
            )

        rows: list[InputCompanyRow] = []
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            company = self._get_cell_value(row, company_column)
            registration_number = ""
            if registration_column is not None:
                registration_number = self._get_cell_value(row, registration_column)

            if not company and not registration_number:
                continue

            rows.append(
                InputCompanyRow(
                    row_index=row_index,
                    company=company,
                    registration_number=registration_number,
                )
            )

        workbook.close()
        return rows

    @staticmethod
    def _to_text(value: object) -> str:
        """Convert a cell value to a normalized string."""
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _build_header_map(header_row: tuple[object, ...]) -> dict[str, int]:
        """Build a mapping from header name to column index."""
        header_map: dict[str, int] = {}
        for index, value in enumerate(header_row):
            header = str(value).strip() if value is not None else ""
            if header:
                header_map[header] = index
        return header_map

    @staticmethod
    def _find_column_index(
        header_map: dict[str, int],
        aliases: tuple[str, ...],
    ) -> int | None:
        """Find the first matching column index by alias."""
        for alias in aliases:
            if alias in header_map:
                return header_map[alias]
        return None

    def _get_cell_value(self, row: tuple[object, ...], column_index: int) -> str:
        """Return a normalized cell value by column index."""
        if column_index >= len(row):
            return ""
        return self._to_text(row[column_index])
